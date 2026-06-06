"""DOCX export (python-docx) — executive summary, readable structured sections,
all DONE agents, and a Sources section. Cost is never included."""
import io
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from .common import (today, done, flatten_output, executive_summary, collect_sources,
                     ORANGE, ORANGE_DARK, INK3)


def _heading(doc, text, size, hex_color):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(hex_color)
    return p


def export(project, executions) -> bytes:
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"Stratégie — {project.name}")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor.from_string(ORANGE)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mrun = meta.add_run(f"Date : {today()} — Strategor")
    mrun.font.size = Pt(11)
    mrun.font.color.rgb = RGBColor.from_string(INK3)

    summary = executive_summary(executions)
    if summary:
        _heading(doc, "Synthèse exécutive", 16, ORANGE_DARK)
        doc.add_paragraph(summary)

    for e in done(executions):
        _heading(doc, f"Agent {e.agentId} — {e.agentName or ''}", 16, ORANGE_DARK)
        blocks = flatten_output(e.output)
        if not blocks:
            doc.add_paragraph("(Aucune donnée)")
        for heading, lines in blocks:
            _heading(doc, heading, 12, INK3)
            if len(lines) == 1 and not lines[0].startswith("• "):
                doc.add_paragraph(lines[0])
            else:
                for ln in lines:
                    doc.add_paragraph(ln.lstrip("• "), style="List Bullet")

    sources = collect_sources(executions)
    if sources:
        _heading(doc, "Sources", 16, ORANGE_DARK)
        for s in sources:
            doc.add_paragraph(s, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

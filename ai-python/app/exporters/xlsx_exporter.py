"""XLSX export (openpyxl) — Synthèse (+ executive summary), Agents (model/tokens,
NO cost), readable Analyses, and a Sources sheet. Cost is never included."""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .common import today, done, flatten_output, executive_summary, collect_sources

HEADER_FILL = PatternFill("solid", fgColor="E8621A")  # brand orange
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _header(ws, row, values):
    for col, val in enumerate(values, start=1):
        c = ws.cell(row=row, column=col, value=val)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT


def export(project, executions) -> bytes:
    wb = Workbook()

    # Synthèse
    s = wb.active
    s.title = "Synthèse"
    _header(s, 1, ["Projet", project.name])
    s.cell(row=2, column=1, value="Date"); s.cell(row=2, column=2, value=today())
    s.cell(row=3, column=1, value="Mode"); s.cell(row=3, column=2, value=project.analysisMode or "comprehensive")
    summary = executive_summary(executions)
    if summary:
        s.cell(row=5, column=1, value="Synthèse exécutive").font = Font(bold=True)
        c = s.cell(row=6, column=1, value=summary)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        s.merge_cells(start_row=6, start_column=1, end_row=6, end_column=6)
    s.column_dimensions["A"].width = 22
    s.column_dimensions["B"].width = 60

    # Agents (all executions) — model + tokens only, NO cost
    a = wb.create_sheet("Agents")
    _header(a, 1, ["ID", "Agent", "Statut", "Modèle", "Tokens IN", "Tokens OUT"])
    r = 2
    for e in executions:
        a.cell(row=r, column=1, value=e.agentId)
        a.cell(row=r, column=2, value=e.agentName or "")
        a.cell(row=r, column=3, value=e.status)
        a.cell(row=r, column=4, value=e.modelUsed or "")
        a.cell(row=r, column=5, value=e.tokensInput or 0)
        a.cell(row=r, column=6, value=e.tokensOutput or 0)
        r += 1
    for col in "ABCDEF":
        a.column_dimensions[col].width = 16

    # Analyses (DONE only) — readable, section by section
    o = wb.create_sheet("Analyses")
    _header(o, 1, ["Agent", "Section", "Contenu"])
    r = 2
    for e in done(executions):
        label = f"{e.agentId} — {e.agentName}"
        blocks = flatten_output(e.output)
        if not blocks:
            o.cell(row=r, column=1, value=label)
            o.cell(row=r, column=3, value="(Aucune donnée)")
            r += 1
            continue
        for heading, lines in blocks:
            o.cell(row=r, column=1, value=label)
            o.cell(row=r, column=2, value=heading)
            txt = "\n".join(ln.lstrip("• ") for ln in lines)
            if len(txt) > 30000:
                txt = txt[:30000] + "…"
            c = o.cell(row=r, column=3, value=txt)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
    o.column_dimensions["A"].width = 28
    o.column_dimensions["B"].width = 24
    o.column_dimensions["C"].width = 90

    # Sources
    sources = collect_sources(executions)
    if sources:
        sh = wb.create_sheet("Sources")
        _header(sh, 1, ["Source"])
        for i, src in enumerate(sources, start=2):
            sh.cell(row=i, column=1, value=src)
        sh.column_dimensions["A"].width = 100

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

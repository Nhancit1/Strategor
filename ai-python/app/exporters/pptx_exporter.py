"""PPTX export (python-pptx) — 16:9 deck: title, executive-summary slide, one
slide per DONE agent (readable structured text), and a Sources slide. No cost."""
import io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from .common import (today, done, flatten_output, executive_summary, collect_sources,
                     ORANGE, ORANGE_DARK, INK, INK3)


def _box(slide, text, left, top, width, height, size, bold, hex_color):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(hex_color)


def _body_text(output) -> str:
    out = []
    for heading, lines in flatten_output(output):
        out.append(heading)
        out.extend(f"  {ln}" if ln.startswith("• ") else f"  • {ln}" for ln in lines)
        out.append("")
    return "\n".join(out).strip() or "(Aucune donnée)"


def export(project, executions) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)   # 16:9 widescreen
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Title slide
    s0 = prs.slides.add_slide(blank)
    _box(s0, "Stratégie", 0.7, 2.7, 12, 1.1, 36, True, ORANGE)
    _box(s0, project.name, 0.7, 3.9, 12, 0.8, 24, False, INK)
    _box(s0, f"{today()} — Strategor", 0.7, 6.5, 12, 0.5, 12, False, INK3)

    # Executive summary slide
    summary = executive_summary(executions)
    if summary:
        s = prs.slides.add_slide(blank)
        _box(s, "Synthèse exécutive", 0.4, 0.4, 12.5, 0.7, 22, True, ORANGE_DARK)
        body = summary if len(summary) <= 1600 else summary[:1600] + "…"
        _box(s, body, 0.4, 1.25, 12.5, 5.9, 13, False, INK)

    # One slide per DONE agent
    for e in done(executions):
        s = prs.slides.add_slide(blank)
        _box(s, f"Agent {e.agentId} — {e.agentName}", 0.4, 0.4, 12.5, 0.7, 20, True, ORANGE_DARK)
        body = _body_text(e.output)
        if len(body) > 1600:
            body = body[:1600] + "…"
        _box(s, body, 0.4, 1.25, 12.5, 5.9, 11, False, INK)

    # Sources slide
    sources = collect_sources(executions)
    if sources:
        s = prs.slides.add_slide(blank)
        _box(s, "Sources", 0.4, 0.4, 12.5, 0.7, 20, True, ORANGE_DARK)
        body = "\n".join(f"• {x}" for x in sources)
        if len(body) > 1600:
            body = body[:1600] + "…"
        _box(s, body, 0.4, 1.25, 12.5, 5.9, 11, False, INK)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()

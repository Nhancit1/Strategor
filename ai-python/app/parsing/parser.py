"""
Document text extraction — port of DocumentParserService.

docx / xlsx / pptx via python-docx / openpyxl / python-pptx.
txt / md / csv read directly. PDF via pdfplumber (the Java version only
stubbed PDFs; here we extract real text).
"""
import logging
import os
from pathlib import Path

from ..config import settings

log = logging.getLogger("strategor.parser")

# ── Resource bounds: stop a small malicious file (zip-bomb / huge sheet / endless PDF)
#    from exploding memory/CPU. Generous for real business documents, fatal for bombs. ──
MAX_FILE_BYTES = 30 * 1024 * 1024   # 30 MB on disk (matches the upload ceiling + margin)
MAX_TEXT_CHARS = 2_000_000          # cap extracted text (~2 MB) fed downstream
MAX_PDF_PAGES = 500
MAX_XLSX_CELLS = 500_000
MAX_PPTX_SLIDES = 500


def _safe_path(storage_path: str) -> str:
    """Resolve storage_path and confine it under the allowed upload root.

    Defeats path traversal / arbitrary file read: a caller cannot make the parser open
    /etc/passwd, secrets, or anything outside the uploads volume (symlinks are resolved
    before the check, so a symlink pointing outside the root is also refused).
    """
    root = Path(settings.storage_root).resolve()
    candidate = Path(storage_path)
    if not candidate.is_absolute():
        candidate = root / candidate
    resolved = candidate.resolve()
    if not resolved.is_relative_to(root):
        raise ValueError("storagePath is outside the allowed upload directory")
    if not resolved.is_file():
        raise FileNotFoundError(f"File not found in upload directory: {storage_path}")
    if resolved.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("File too large to parse")
    return str(resolved)


def extract_text(storage_path: str, filename: str, mime_type: str | None) -> str:
    storage_path = _safe_path(storage_path)  # confine + size-check BEFORE any open()
    name = (filename or "").lower()
    mime = mime_type or ""

    if name.endswith(".docx") or "wordprocessing" in mime:
        text = _docx(storage_path)
    elif name.endswith(".xlsx") or "spreadsheet" in mime:
        text = _xlsx(storage_path)
    elif name.endswith(".pptx") or "presentation" in mime:
        text = _pptx(storage_path)
    elif name.endswith((".txt", ".md", ".csv")) or mime.startswith("text/"):
        with open(storage_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read(MAX_TEXT_CHARS + 1)
    elif name.endswith(".pdf") or "pdf" in mime:
        text = _pdf(storage_path)
    else:
        text = (f"[Document {filename} — type non supporté pour extraction. "
                "Contenu à exploiter manuellement.]")

    # Final guard: never return more than the cap, whatever the extractor produced.
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS] + "\n…[contenu tronqué — limite de sécurité]…"
    return text


def _docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    lines = [p.text for p in doc.paragraphs]
    # tables too
    for table in doc.tables:
        for row in table.rows:
            lines.append("\t".join(c.text for c in row.cells))
    return "\n".join(lines)


def _xlsx(path: str) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    seen = 0  # total cells visited — hard stop guards against a tiny zip-bomb sheet
    for ws in wb.worksheets:
        out.append(f"== Feuille : {ws.title} ==")
        for row in ws.iter_rows(values_only=True):
            seen += len(row)
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                out.append("\t".join(cells))
            if seen >= MAX_XLSX_CELLS:
                out.append("…[feuille tronquée — limite de sécurité]…")
                wb.close()
                return "\n".join(out)
    wb.close()
    return "\n".join(out)


def _pptx(path: str) -> str:
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        if i > MAX_PPTX_SLIDES:
            out.append("…[présentation tronquée — limite de sécurité]…")
            break
        out.append(f"== Slide {i} ==")
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
    return "\n".join(out)


def _pdf(path: str) -> str:
    import pdfplumber
    out = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            if i > MAX_PDF_PAGES:
                out.append("…[PDF tronqué — limite de sécurité]…")
                break
            text = page.extract_text() or ""
            out.append(f"== Page {i} ==\n{text}")
    return "\n".join(out)
